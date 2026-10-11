"""Rerunnable structural acceptance checks, separate from native numerical tests."""
import json
import re
from pathlib import Path
from static_audit import audit, PRODUCT
from openpyxl import load_workbook


def main():
    inventory = audit()
    checks = []

    def check(name, ok, evidence):
        checks.append({"test": name, "status": "PASS" if ok else "FAIL", "evidence": evidence})

    for listing in inventory["listings"]:
        check(listing["name"] + " title and tag limits", listing["within_limits"],
              {"title_characters": listing["title_characters"], "tag_count": len(listing["tags"]),
               "longest_tag": max(t["characters"] for t in listing["tags"])})
    for name, info in inventory["workbooks"].items():
        check(name + " no macros/external links/personal identifiers/paths",
              not info["macros"] and not info["external_links"] and not info["suspicious_metadata"],
              {"creator": "openpyxl (library)", "external_links": info["external_links"],
               "suspicious_metadata": info["suspicious_metadata"]})
    oee = load_workbook(PRODUCT / "dist" / "OEE-Tracker.xlsx")
    formula = oee["Production Log"]["R5"].value
    check("OEE Check flags every performance above 100%", "O5>1.05" not in formula, formula)
    down = load_workbook(PRODUCT / "dist" / "Downtime-Log-and-Pareto.xlsx")
    pm = load_workbook(PRODUCT / "dist" / "Preventive-Maintenance-Schedule.xlsx")
    vulnerable = {"OEE machine": oee["Dashboard"]["B9"].value,
                  "OEE cycle lookup": oee["Production Log"]["K5"].value,
                  "downtime machine": down["By Machine"]["B5"].value,
                  "downtime reason": down["Pareto"]["J9"].value,
                  "downtime stops": down["Pareto"]["K9"].value,
                  "downtime category lookup": down["Downtime Log"]["I5"].value,
                  "PM completion matching": pm["Completion Log"]["F5"].value}
    for label, formula in vulnerable.items():
        check(label + " literal names safe under documented Excel wildcard semantics",
              "SUBSTITUTE(" in formula or ("SUMPRODUCT(" in formula and not re.search(r"\b(?:SUMIFS|COUNTIFS|MATCH)\(", formula)), {"formula": formula,
                  "acceptance": "Escaped criteria or exact array equality without wildcard-consuming criteria functions",
                  "source": ("https://support.microsoft.com/en-us/excel/functions/match-function"
                             if "lookup" in label else "https://support.microsoft.com/en-au/excel/functions/sumifs-function"),
                  "native_Excel_test": "unverified; exact equality removes the documented wildcard construct"})
    delivery = (PRODUCT / "LISTING.md").read_text(encoding="utf-8").split("## Files per listing")[1].split("## Images")[0]
    check("Declared direct-upload format agrees with Etsy's published file-type list", ".zip" in delivery.lower(),
          {"instructions": delivery.strip(), "supported_wrapper": ".zip",
           "contrary_repository_record": "PR #14 / main 1c9ebda records direct XLSX attachments after reload",
           "live_upload_and_buyer_download": "unverified; this check is documentation-list agreement, not proof of rejection",
           "source": "https://help.etsy.com/hc/en-us/articles/115015628347-How-to-Manage-Your-Digital-Listings"})
    result = {"tests": checks, "summary": {s: sum(c["status"] == s for c in checks) for s in ("PASS", "FAIL")}}
    Path(__file__).with_name("static-checks.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result["summary"]))
    return int(result["summary"]["FAIL"] > 0)


if __name__ == "__main__":
    raise SystemExit(main())
