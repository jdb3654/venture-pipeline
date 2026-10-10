"""Adversarial acceptance suite. Real soffice results; no formula emulation.

PASS/FAIL are product acceptance results. Infrastructure failures are BLOCKED.
Known defects deliberately remain failures. All workbook writes are temporary.
"""
import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import re
import socket
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import date, datetime, time as clock_time

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

HERE = Path(__file__).resolve().parent
PRODUCT = HERE.parent
NAMES = {"oee": "OEE-Tracker.xlsx", "down": "Downtime-Log-and-Pareto.xlsx",
         "pm": "Preventive-Maintenance-Schedule.xlsx"}
EPOCH = date(1899, 12, 30)
BASE_DATE = date(2026, 9, 1)


def serial(d):
    return (d - EPOCH).days


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def blank_inputs(wb, sheet, cols, last):
    for row in range(5, last + 1):
        for col in cols:
            wb[sheet].cell(row, col).value = None


def fresh(kind):
    wb = load_workbook(PRODUCT / "dist" / NAMES[kind])
    if kind == "oee":
        blank_inputs(wb, "Production Log", range(1, 9), 504)
        blank_inputs(wb, "Setup", range(1, 4), 24)
    elif kind == "down":
        blank_inputs(wb, "Downtime Log", [1, 2, 3, 4, 5, 6, 8, 10], 1004)
        blank_inputs(wb, "Setup", [1, 4, 5], 34)
    else:
        blank_inputs(wb, "PM Plan", [1, 2, 3, 4, 5, 10], 304)
        blank_inputs(wb, "Completion Log", range(1, 6), 2004)
    return wb


def metrics(rows):
    """Independent aggregation of raw (shift, stops, down, parts, rejects, cycle)."""
    planned = sum(r[0] - r[1] for r in rows)
    run = sum(r[0] - r[1] - r[2] for r in rows)
    parts = sum(r[3] for r in rows)
    good = sum(r[3] - r[4] for r in rows)
    ideal = sum(r[5] * r[3] / 60 for r in rows)
    # Quality factor used in the delivered dashboard.
    factors = [run / planned, ideal / run, good / parts]
    return [planned, run, sum(r[2] for r in rows), parts, good, ideal,
            *factors, math.prod(factors)]


def oee_row(wb, row=5, machine="M1", shift=480, stops=40, down=35, parts=1900,
            rejects=22, day=BASE_DATE):
    for col, value in enumerate([day, "Day", machine, shift, stops, down, parts, rejects], 1):
        wb["Production Log"].cell(row, col, value)


def down_row(wb, row=5, machine="M1", reason="R1", minutes=10,
             start=None, end=None, day=BASE_DATE):
    for col, value in zip([1, 2, 3, 4, 5, 6, 8],
                          [day, "Day", machine, start, end, minutes, reason]):
        wb["Downtime Log"].cell(row, col).value = value


def pm_row(wb, row=5, asset="A1", task="T1", frequency=7, earlier=None):
    for col, value in zip([1, 2, 3, 5], [asset, task, frequency, earlier]):
        wb["PM Plan"].cell(row, col).value = value


def heterogeneous_input():
    wb = fresh("oee")
    for row, machine, cycle in [(5, "Slow", 60), (6, "Fast", 6)]:
        wb["Setup"].cell(row, 1, machine)
        wb["Setup"].cell(row, 2, cycle)
    oee_row(wb, machine="Slow", shift=60, stops=0, down=0, parts=60, rejects=0)
    oee_row(wb, row=6, machine="Fast", shift=60, stops=0, down=0, parts=600, rejects=600)
    return wb


class Suite:
    def __init__(self, root, reuse=False):
        self.root, self.requests, self.checks, self.reuse = root, [], [], reuse

    def scenario(self, id, kind, wb, read=None, features=None, pdf=None, original=False):
        path = self.root / (id + ".xlsx")
        if not self.reuse or not path.exists():
            if original:
                shutil.copyfile(PRODUCT / "dist" / NAMES[kind], path)
            else:
                wb.save(path)
        self.requests.append({"id": id, "input": str(path), "read": read or {},
                              "features": features or {}, "pdf": pdf})
        return id

    def check(self, id, title, fn):
        self.checks.append((id, title, fn))

    def equal(self, id, title, sheet, cell, expected, tolerance=1e-8):
        def check(result):
            actual = get(result, sheet, cell)
            ok = (isinstance(actual, (int, float)) and abs(actual - expected) <= tolerance
                  if isinstance(expected, (int, float)) else actual == expected)
            return ok, {"cell": sheet + "!" + cell, "expected": expected, "actual": actual}
        self.check(id, title, check)

    def non_ok(self, id, title, sheet, cell):
        self.check(id, title, lambda r: (isinstance(get(r, sheet, cell), str) and
                                        get(r, sheet, cell) not in ("OK", "", "Due soon", "OVERDUE"),
                                        {"cell": sheet + "!" + cell, "actual": get(r, sheet, cell),
                                         "expected": "explicit invalid/missing-input warning"}))


def get(result, sheet, cell):
    col, row, _, _ = range_boundaries(cell)
    for rng, data in result["data"][sheet].items():
        c1, r1, c2, r2 = range_boundaries(rng)
        if c1 <= col <= c2 and r1 <= row <= r2:
            return data[row - r1][col - c1]
    raise KeyError(sheet + "!" + cell)


def scenarios(s):
    # Shipped examples, every shift and machine, both weighted and heterogeneous totals.
    wb = load_workbook(PRODUCT / "dist" / NAMES["oee"])
    reads = {"Production Log": ["A5:R504"], "Dashboard": ["A9:L29"]}
    id = s.scenario("oee_examples", "oee", wb, reads,
        {"Setup": ["A5"], "Production Log": ["C5", "D5", "Q5"], "Dashboard": ["K9"]},
        {"Dashboard": "A1:L48"})
    grouped = {}
    for row in range(5, 11):
        raw = [wb["Production Log"].cell(row, c).value for c in range(4, 9)]
        name = wb["Production Log"].cell(row, 3).value
        cycle = {wb["Setup"].cell(r, 1).value: wb["Setup"].cell(r, 2).value for r in range(5, 8)}[name]
        m = metrics([raw + [cycle]])
        grouped.setdefault(name, []).append(raw + [cycle])
        for col, expected in zip("NOPQ", m[6:]):
            s.equal(id, "example shift %s %s" % (row, col), "Production Log", col + str(row), expected)
    all_rows = []
    for row, (name, rows) in enumerate(grouped.items(), 9):
        all_rows += rows
        for col, expected in zip("BCDEFGHIJK", metrics(rows)):
            s.equal(id, "machine %s %s" % (name, col), "Dashboard", col + str(row), expected)
    for col, expected in zip("BCDEFGHIJK", metrics(all_rows)):
        s.equal(id, "all machines summed-factor %s" % col, "Dashboard", col + "29", expected)
    # Correct net productive time aggregation accounts for each row's cycle and good units.
    productive = sum(r[5] * (r[3] - r[4]) / 60 for r in all_rows)
    s.equal(id, "heterogeneous machines OEE = sum ideal good time / sum planned time",
            "Dashboard", "K29", productive / sum(r[0] - r[1] for r in all_rows))
    s.check(id, "equal planned example shifts agree with their average (weighting tested separately)", lambda r: (
        abs(get(r, "Dashboard", "K9") - sum(get(r, "Production Log", "Q%d" % x)
            for x in (5, 8)) / 2) < 1e-8,
        "Examples have equal planned times; unequal-shift scenario establishes weighting, not this equality."))

    id = s.scenario("oee_heterogeneous", "oee", heterogeneous_input(), reads)
    # Slow: 60 ideal good minutes. Fast: zero ideal good minutes. 120 planned.
    s.equal(id, "different cycle/reject mix yields 60/120 = 50%, not 60/660 = 9.09%",
            "Dashboard", "K29", .5)

    for name, start, end in [("blank", None, None), ("from", date(2026, 9, 2), None),
                             ("to", None, BASE_DATE), ("both", BASE_DATE, date(2026, 9, 2)),
                             ("boundary", BASE_DATE, BASE_DATE),
                             ("outside", date(2026, 8, 1), date(2026, 8, 2))]:
        wb = load_workbook(PRODUCT / "dist" / NAMES["oee"])
        wb["Dashboard"]["B3"], wb["Dashboard"]["B4"] = start, end
        id = s.scenario("oee_date_" + name, "oee", wb, reads)
        selected = []
        for row in range(5, 11):
            day = wb["Production Log"].cell(row, 1).value.date()
            if (start is None or day >= start) and (end is None or day <= end):
                selected.append(row)
        s.equal(id, "date filter %s inclusive planned time" % name, "Dashboard", "B29", 440 * len(selected))

    for name, kwargs in [("rejects", {"rejects": 2000}), ("downtime", {"down": 450}),
        ("planned_stops", {"stops": 500}), ("zero_parts", {"parts": 0, "rejects": 0}),
        ("zero_shift", {"shift": 0, "stops": 0, "down": 0, "parts": 0, "rejects": 0}),
        ("unknown_machine", {"machine": "TYPO"}), ("missing_cycle", {}),
        ("performance_102", {"parts": 2065, "rejects": 0}),
        ("performance_110", {"parts": 2250, "rejects": 0}),
        ("negative_parts", {"parts": -1, "rejects": 0}),
        ("text_parts", {"parts": "one thousand"})]:
        wb = fresh("oee")
        wb["Setup"]["A5"], wb["Setup"]["B5"] = "M1", (None if name == "missing_cycle" else 12)
        oee_row(wb, **kwargs)
        id = s.scenario("oee_bad_" + name, "oee", wb, reads)
        s.non_ok(id, "Check catches " + name, "Production Log", "R5")
        s.equal(id, "invalid %s suppresses OEE" % name, "Production Log", "Q5", "")

    wb = fresh("oee")
    wb["Setup"]["A5"], wb["Setup"]["B5"] = "M1", 60
    oee_row(wb, shift=60, stops=0, down=0, parts=30, rejects=0)
    oee_row(wb, row=6, shift=240, stops=0, down=0, parts=240, rejects=0)
    id = s.scenario("oee_unequal", "oee", wb, reads)
    s.equal(id, "unequal shifts time-weighted 270/300, not average 75%", "Dashboard", "K9", .9)

    for target in (.95, .5):
        wb = load_workbook(PRODUCT / "dist" / NAMES["oee"])
        wb["Dashboard"]["B5"] = target
        id = s.scenario("oee_target_" + str(target), "oee", wb, reads,
                        {"Production Log": ["Q5"], "Dashboard": ["K9"]},
                        {"Dashboard": "A1:L48"})
        s.check(id, "vs target follows %.0f%%" % (target * 100), lambda r, t=target: (
            abs(get(r, "Dashboard", "L9") - (get(r, "Dashboard", "K9") - t)) < 1e-8,
            {"K9": get(r, "Dashboard", "K9"), "L9": get(r, "Dashboard", "L9"), "target": t}))

    wb = fresh("oee")
    cap_rows = []
    for i in range(20):
        wb["Setup"].cell(5 + i, 1, "M%02d" % i)
        wb["Setup"].cell(5 + i, 2, 30)
    for i in range(500):
        oee_row(wb, row=5 + i, machine="M%02d" % (i % 20), shift=480, stops=30,
                down=i % 45, parts=600 + i % 100, rejects=i % 10)
        cap_rows.append([480, 30, i % 45, 600 + i % 100, i % 10, 30])
    id = s.scenario("oee_capacity", "oee", wb, reads)
    for col, expected in zip("BCDEFGHIJK", metrics(cap_rows)):
        s.equal(id, "500 rows 20 machines %s" % col, "Dashboard", col + "29", expected)
    s.equal(id, "last OEE row counted", "Production Log", "R504", "OK")

    down_reads = {"Downtime Log": ["A5:J1004"], "Pareto": ["A6:F38", "I9:L38"],
                  "By Machine": ["A5:E34"]}
    wb = load_workbook(PRODUCT / "dist" / NAMES["down"])
    id = s.scenario("down_examples", "down", wb, down_reads,
        {"Setup": ["A5"], "Downtime Log": ["C5", "H5"], "Pareto": ["B5"], "By Machine": ["A5"]},
        {"Pareto": "A1:Z24"})
    for row, expected in enumerate([35, 15, 60, 45, 55, 30, 15], 5):
        s.equal(id, "example duration %d" % row, "Downtime Log", "G%d" % row, expected)
    s.equal(id, "blank machine filter includes all 255 minutes", "Pareto", "B6", 255)
    for row, expected in enumerate([105, 75, 75], 5):
        s.equal(id, "By Machine example %d" % row, "By Machine", "B%d" % row, expected)
    s.check(id, "By Machine reconciles to Pareto unfiltered", lambda r: (
        sum(get(r, "By Machine", "B%d" % i) for i in (5, 6, 7)) == get(r, "Pareto", "B6"),
        {"By Machine": sum(get(r, "By Machine", "B%d" % i) for i in (5, 6, 7)),
         "Pareto": get(r, "Pareto", "B6")}))

    wb = fresh("down")
    wb["Setup"]["A5"], wb["Setup"]["D5"], wb["Setup"]["E5"] = "M1", "R1", "Cat"
    durations = [("normal", clock_time(7, 15), clock_time(7, 50), None, 35),
        ("overnight", clock_time(23, 40), clock_time(0, 25), None, 45),
        ("equal", clock_time(12), clock_time(12), None, 0),
        ("midnight_end", clock_time(23, 40), clock_time(0), None, 20),
        ("midnight_start", clock_time(0), clock_time(0, 25), None, 25),
        ("override", clock_time(7), clock_time(8), 7, 7),
        ("zero_override", clock_time(7), clock_time(8), 0, 0),
        ("missing_start", None, clock_time(8), None, ""),
        ("missing_end", clock_time(7), None, None, ""),
        ("missing_both", None, None, None, "")]
    for row, (_, a, b, minutes, _) in enumerate(durations, 5):
        down_row(wb, row=row, start=a, end=b, minutes=minutes)
    wb["Pareto"]["B5"] = "M1"
    id = s.scenario("down_durations", "down", wb, down_reads)
    for row, (name, _, _, _, expected) in enumerate(durations, 5):
        s.equal(id, "duration " + name, "Downtime Log", "G%d" % row, expected)
    s.check(id, "stop counts use same definition on Pareto and By Machine", lambda r: (
        get(r, "Pareto", "D9") == get(r, "By Machine", "C5"),
        {"Pareto": get(r, "Pareto", "D9"), "By Machine": get(r, "By Machine", "C5")}))

    for kind, literals in [("machine", ["Line*1", "Line?1", "Line~1"]),
                           ("reason", ["Jam*1", "Jam?1", "Jam~1"]),
                           ("regex_reason", ["R.1", "R(1)", "R+1"])]:
        for literal in literals:
            wb = fresh("down")
            counterpart = literal.replace("*", "X").replace("?", "X").replace("~", "")
            if counterpart == literal:
                counterpart = {"R.1": "RX1", "R(1)": "R1", "R+1": "RR1"}[literal]
            for row, name in enumerate([literal, counterpart], 5):
                wb["Setup"].cell(row, 1, name if kind == "machine" else "M1")
                wb["Setup"].cell(row, 4, "R1" if kind == "machine" else name)
                wb["Setup"].cell(row, 5, "Cat")
                down_row(wb, row=row, machine=name if kind == "machine" else "M1",
                         reason="R1" if kind == "machine" else name, minutes=10 if row == 5 else 90)
            if kind != "machine":
                wb["Setup"]["A6"] = None
            wb["Pareto"]["B5"] = literal if kind == "machine" else "M1"
            id = s.scenario("down_literal_%s_%s" % (kind, literals.index(literal)), "down", wb, down_reads)
            s.equal(id, "literal %s %s minutes" % (kind, literal),
                    "By Machine" if kind == "machine" else "Pareto",
                    "B5" if kind == "machine" else "J9", 10)
            if kind == "machine":
                s.equal(id, "literal machine filter " + literal, "Pareto", "B6", 10)
            else:
                s.equal(id, "literal reason total no duplication " + literal, "Pareto", "B6", 100)
            s.equal(id, "literal reason category lookup " + literal, "Downtime Log", "I5", "Cat")

    for mode in ("ranking", "unknown", "blank_machine", "combined", "capacity", "negative"):
        wb = fresh("down")
        wb["Setup"]["A5"], wb["Setup"]["A6"] = "M1", "M2"
        for i in range(30):
            wb["Setup"].cell(5 + i, 4, "R%02d" % i)
            wb["Setup"].cell(5 + i, 5, "Cat")
        if mode == "ranking":
            minutes = [100, 100] + list(range(28, 0, -1))
            minutes[-1] = 0
            for i, m in enumerate(minutes):
                down_row(wb, row=5 + i, reason="R%02d" % i, minutes=m)
            wb["Pareto"]["B5"] = "M1"
        elif mode == "capacity":
            minutes = [1 + i % 60 for i in range(1000)]
            for i, m in enumerate(minutes):
                down_row(wb, row=5 + i, machine="M%d" % (1 + i % 2), reason="R%02d" % (i % 30), minutes=m)
        elif mode == "unknown":
            down_row(wb, reason="TYPO", minutes=90)
            down_row(wb, row=6, reason="R00", minutes=10)
            wb["Pareto"]["B5"] = "M1"
        elif mode == "blank_machine":
            down_row(wb, machine=None, reason="R00", minutes=90)
            down_row(wb, row=6, reason="R00", minutes=10)
        elif mode == "combined":
            for i, (machine, day, m) in enumerate([
                ("M1", BASE_DATE, 10), ("M2", BASE_DATE, 90),
                ("M1", date(2026, 9, 2), 20), ("M2", date(2026, 9, 2), 80),
                ("M1", date(2026, 9, 3), 30)]):
                down_row(wb, row=5 + i, machine=machine, reason="R00", minutes=m, day=day)
            wb["Pareto"]["B3"], wb["Pareto"]["B4"], wb["Pareto"]["B5"] = BASE_DATE, date(2026, 9, 2), "M1"
        else:
            down_row(wb, reason="R00", minutes=-20)
            down_row(wb, row=6, reason="R01", minutes=10)
            wb["Pareto"]["B5"] = "M1"
        id = s.scenario("down_" + mode, "down", wb, down_reads,
                        {"Pareto": ["B5"]} if mode == "ranking" else None,
                        {"Pareto": "A1:Z24"} if mode == "ranking" else None)
        if mode == "ranking":
            s.equal(id, "all 30 reason minutes table total", "Pareto", "B6", sum(minutes))
            s.equal(id, "ties stable Setup order first", "Pareto", "B9", "R00")
            s.equal(id, "ties stable Setup order second", "Pareto", "B10", "R01")
            s.equal(id, "zero minute reason hidden", "Pareto", "B38", "")
            s.equal(id, "positive reason cumulative ends at 100%", "Pareto", "F37", 1)
            s.equal(id, "top 12 cumulative uses all reasons denominator", "Pareto", "F20", sum(minutes[:12]) / sum(minutes))
            for i, m in enumerate(minutes[:-1]):
                s.equal(id, "ranked reason %02d minutes" % i, "Pareto", "C%d" % (9 + i), m)
        elif mode == "capacity":
            s.equal(id, "1000 rows 30 reasons total", "Pareto", "B6", sum(minutes))
            s.equal(id, "capacity cumulative 100%", "Pareto", "F38", 1)
            s.equal(id, "1000th duration counted", "Downtime Log", "G1004", minutes[-1])
            s.check(id, "capacity By Machine total reconciles", lambda r: (
                sum(get(r, "By Machine", "B%d" % i) for i in (5, 6)) == sum(minutes),
                {"expected": sum(minutes), "actual": sum(get(r, "By Machine", "B%d" % i) for i in (5, 6))}))
        elif mode == "unknown":
            s.equal(id, "unknown reason warning shown", "Downtime Log", "I5", "Not in Setup")
            s.equal(id, "unknown reason not lost from total downtime", "Pareto", "B6", 100)
            s.equal(id, "By Machine includes unknown reason", "By Machine", "B5", 100)
        elif mode == "blank_machine":
            s.equal(id, "blank all-machine filter includes blank-machine downtime", "Pareto", "B6", 100)
        elif mode == "combined":
            s.equal(id, "date + machine intersection inclusive", "Pareto", "B6", 30)
            s.equal(id, "By Machine ignores machine filter but follows dates", "By Machine", "B6", 170)
        else:
            s.check(id, "negative duration rejected or warned", lambda r: (
                not isinstance(get(r, "Downtime Log", "G5"), (int, float))
                or get(r, "Downtime Log", "G5") >= 0,
                {"G5": get(r, "Downtime Log", "G5"), "Pareto total": get(r, "Pareto", "B6")}))

    pm_reads = {"PM Plan": ["A1:L304"], "Completion Log": ["A5:F2004"]}
    wb = load_workbook(PRODUCT / "dist" / NAMES["pm"])
    id = s.scenario("pm_examples", "pm", wb, pm_reads,
        {"PM Plan": ["I5", "I6", "I9"], "Completion Log": ["B5", "C5"]},
        {"PM Plan": "A1:L11"})
    s.check(id, "TODAY examples remain relative on actual office date", lambda r: (
        [get(r, "PM Plan", "H%d" % i) for i in range(5, 10)] == [-10, 4, 3, 12, 160],
        {"days left": [get(r, "PM Plan", "H%d" % i) for i in range(5, 10)]}))
    s.equal(id, "example overdue count", "PM Plan", "L1", 1)
    s.equal(id, "example due-soon count", "PM Plan", "L2", 2)
    s.equal(id, "example OK count", "PM Plan", "L3", 2)

    for mode in ("never", "latest", "earlier_later", "typo_asset", "typo_task", "blank_frequency",
                 "zero_frequency", "negative_frequency", "window0", "window30", "duplicate",
                 "wildcard", "blank_date"):
        wb = fresh("pm")
        pm_row(wb, frequency=None if mode == "blank_frequency" else 0 if mode == "zero_frequency"
               else -7 if mode == "negative_frequency" else 7,
               earlier=date(2026, 9, 30) if mode == "earlier_later" else None)
        if mode not in ("never", "blank_frequency", "zero_frequency", "negative_frequency"):
            wb["Completion Log"]["A5"] = date(2026, 9, 1)
            wb["Completion Log"]["B5"] = "TYPO" if mode == "typo_asset" else "A*" if mode == "wildcard" else "A1"
            wb["Completion Log"]["C5"] = "TYPO" if mode == "typo_task" else "T1"
        if mode == "latest":
            for row, day in enumerate([date(2026, 9, 9), date(2026, 9, 3)], 6):
                wb["Completion Log"].cell(row, 1, day)
                wb["Completion Log"].cell(row, 2, "A1")
                wb["Completion Log"].cell(row, 3, "T1")
        if mode.startswith("window"):
            wb["Completion Log"]["A5"] = "=TODAY()-3"
            wb["PM Plan"]["I2"] = int(mode[6:])
        if mode == "duplicate":
            pm_row(wb, row=6, frequency=30)
        if mode == "blank_date":
            wb["Completion Log"]["A5"] = None
        id = s.scenario("pm_" + mode, "pm", wb, pm_reads)
        if mode == "never":
            s.equal(id, "never completed days left zero", "PM Plan", "H5", 0)
            s.equal(id, "never completed due today", "PM Plan", "G5", serial(date.today()))
            s.equal(id, "never completed status due soon", "PM Plan", "I5", "Due soon")
        elif mode == "latest":
            s.equal(id, "latest of multiple unsorted completions", "PM Plan", "F5", serial(date(2026, 9, 9)))
        elif mode == "earlier_later":
            s.equal(id, "pre-log later than logged date wins", "PM Plan", "F5", serial(date(2026, 9, 30)))
        elif mode in ("typo_asset", "typo_task", "wildcard"):
            s.equal(id, "completion %s warning" % mode, "Completion Log", "F5", "No matching asset + task in plan")
            s.equal(id, "completion %s leaves plan unchanged" % mode, "PM Plan", "F5", "")
        elif mode in ("blank_frequency", "zero_frequency", "negative_frequency"):
            s.non_ok(id, "invalid %s explicitly warned" % mode, "PM Plan", "I5")
        elif mode.startswith("window"):
            s.equal(id, mode + " follows setting", "PM Plan", "I5", "OK" if mode == "window0" else "Due soon")
        elif mode == "duplicate":
            s.check(id, "duplicate asset+task explicitly warned", lambda r: (
                "duplicate" in str(get(r, "PM Plan", "I5")).lower(),
                {"statuses": [get(r, "PM Plan", "I5"), get(r, "PM Plan", "I6")],
                 "last completed": [get(r, "PM Plan", "F5"), get(r, "PM Plan", "F6")],
                 "next due": [get(r, "PM Plan", "G5"), get(r, "PM Plan", "G6")]}))
        elif mode == "blank_date":
            s.non_ok(id, "undated completion not reported OK", "Completion Log", "F5")
        s.check(id, mode + " status counts match all plan rows", lambda r: (
            [get(r, "PM Plan", "L%d" % x) for x in (1, 2, 3)] ==
            [sum(get(r, "PM Plan", "I%d" % row) == status for row in range(5, 305))
             for status in ("OVERDUE", "Due soon", "OK")], "counts independently counted from plan statuses"))

    # Byte-for-byte baseline avoids blaming openpyxl's temporary input serialization.
    for kind, reads, feature_cells, pdf in [
        ("oee", reads, {"Setup": ["A5"], "Production Log": ["C5", "D5", "Q5"], "Dashboard": ["K9"]}, {"Dashboard": "A1:L48"}),
        ("down", down_reads, {"Setup": ["A5"], "Downtime Log": ["C5", "H5"], "Pareto": ["B5"], "By Machine": ["A5"]}, {"Pareto": "A1:Z24"}),
        ("pm", pm_reads, {"PM Plan": ["I5", "I6", "I9"], "Completion Log": ["B5", "C5"]}, {"PM Plan": "A1:L11"})]:
        id = s.scenario("original_" + kind, kind, None, reads, feature_cells, pdf, original=True)
        expected_freeze = {
            "oee": {"Setup": [True, 0, 4], "Production Log": [True, 3, 4], "Dashboard": [True, 1, 8]},
            "down": {"Setup": [True, 0, 4], "Downtime Log": [True, 3, 4], "Pareto": [True, 0, 8], "By Machine": [True, 0, 4]},
            "pm": {"PM Plan": [True, 2, 4], "Completion Log": [True, 0, 4]}}[kind]
        for sheet, expected in expected_freeze.items():
            s.check(id, sheet + " native frozen position", lambda r, sh=sheet, ex=expected: (
                r["features"][sh]["freeze"] == ex, {"expected": ex, "actual": r["features"][sh]["freeze"]}))
        expected_lists = {"oee": [("Production Log", "C5")],
                          "down": [("Downtime Log", "C5"), ("Downtime Log", "H5"), ("Pareto", "B5")],
                          "pm": [("Completion Log", "B5"), ("Completion Log", "C5")]}[kind]
        for sheet, address in expected_lists:
            s.check(id, sheet + "!" + address + " native list validation", lambda r, sh=sheet, a=address: (
                r["features"][sh]["cells"][a]["validation_type"] == "LIST" and
                bool(r["features"][sh]["cells"][a]["formula1"]), r["features"][sh]["cells"][a]))
        if kind in ("oee", "pm"):
            sheet, address = ("Dashboard", "K9") if kind == "oee" else ("PM Plan", "I5")
            s.check(id, "native three conditional colour styles", lambda r, sh=sheet, a=address: (
                sorted(x["background"] for x in r["features"][sh]["cells"][a]["conditional_rules"]) ==
                sorted([0xC6EFCE, 0xFFE699, 0xF8CBAD]), r["features"][sh]["cells"][a]["conditional_rules"]))
        if kind == "down":
            s.check(id, "native Pareto column plus secondary-axis line", lambda r: (
                chart_axes(r, "Pareto") == [("com.sun.star.chart2.ColumnChartType", 0),
                                            ("com.sun.star.chart2.LineChartType", 1)], chart_axes(r, "Pareto")))
        if kind == "oee":
            s.check(id, "native OEE column chart type", lambda r: (
                chart_axes(r, "Dashboard") == [("com.sun.star.chart2.ColumnChartType", 0)], chart_axes(r, "Dashboard")))
        if kind == "down":
            s.equal(id, "unaltered shipped blank-filter Pareto total", "Pareto", "B6", 255)

    # Unsupported IFERROR otherwise masks deeper OEE defects in this AOO version.
    # Repeat OEE scenarios on explicitly labeled test-only compatibility copies.
    # This is not a fix or a claim that the shipped files work in OpenOffice.
    oee_requests = [r for r in s.requests if r["id"].startswith("oee_")]
    oee_checks = [c for c in s.checks if c[0].startswith("oee_")]
    for req in oee_requests:
        wb = load_workbook(req["input"])
        for ws in wb:
            for row in ws:
                for cell in row:
                    if cell.data_type == "f" and "IFERROR(" in cell.value:
                        cell.value = expand_iferror(cell.value)
        id = s.scenario("diagnostic_" + req["id"], "oee", wb, req["read"], req["features"], req["pdf"])
        for old_id, title, check in oee_checks:
            if old_id == req["id"]:
                s.check(id, "TEST-ONLY IFERROR substitute: " + title, check)

    # Every scenario is scanned, including all preallocated empty formula rows.
    for req in s.requests:
        s.check(req["id"], "zero recalculated formula errors including empty rows", lambda r: (
            not any(r["errors"].values()), r["errors"]))


def expand_iferror(formula):
    """Expand the single IFERROR(expr, \"\") used by build.py; reject other shapes."""
    start = formula.index("IFERROR(")
    pos, depth, quoted = start + len("IFERROR("), 0, False
    comma = None
    for i in range(pos, len(formula)):
        char = formula[i]
        if char == '"':
            quoted = not quoted
        elif not quoted:
            if char == "(":
                depth += 1
            elif char == ")":
                if depth == 0:
                    end = i
                    break
                depth -= 1
            elif char == "," and depth == 0:
                comma = i
    if comma is None or formula[comma + 1:end] != '""':
        raise ValueError("unexpected IFERROR shape")
    expr = formula[pos:comma]
    return formula[:start] + 'IF(ISERROR(%s),"",%s)' % (expr, expr) + formula[end + 1:]


def chart_axes(result, sheet):
    return [(ct["type"], series["axis"])
            for chart in result["features"][sheet]["charts"]
            for system in chart["coordinate_systems"] for ct in system for series in ct["series"]]


def run_office(args, root, requests):
    manifest, output = root / "manifest.json", root / "office-results.json"
    manifest.write_text(json.dumps(requests), encoding="utf-8")
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    profile = root / "office-profile"
    command = [str(args.soffice), "-env:UserInstallation=" + profile.as_uri(), "-headless",
               "-norestore", "-nodefault", "-nofirststartwizard",
               "-accept=socket,host=127.0.0.1,port=%d;urp;StarOffice.ServiceManager" % port]
    creationflags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    log_path = root / "office.log"
    with log_path.open("w") as log:
        office = subprocess.Popen(command, stdout=log, stderr=log, creationflags=creationflags)
        try:
            run = subprocess.run([str(args.uno_python), str(HERE / "office_bridge.py"),
                                  str(manifest), str(output), str(port)],
                                 cwd=str(args.soffice.parent), capture_output=True, text=True,
                                 timeout=args.timeout, creationflags=creationflags)
            (root / "bridge.log").write_text(run.stdout + run.stderr, encoding="utf-8")
            if run.returncode:
                raise RuntimeError(run.stdout + run.stderr)
            return json.loads(output.read_text())
        finally:
            # The UNO bridge terminates only our isolated listener; never kill another user's office.
            try:
                office.wait(timeout=10)
            except subprocess.TimeoutExpired:
                office.terminate()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--soffice", type=Path, default=Path(r"C:\Program Files (x86)\OpenOffice 4\program\soffice.exe"))
    parser.add_argument("--uno-python", type=Path, default=Path(r"C:\Program Files (x86)\OpenOffice 4\program\python.exe"))
    parser.add_argument("--output", type=Path, default=HERE / "results.json")
    parser.add_argument("--work-dir", type=Path)
    parser.add_argument("--only", help="scenario prefix for a bounded retest")
    parser.add_argument("--reuse-inputs", action="store_true", help="reuse prepared copies in --work-dir (diagnostics only)")
    parser.add_argument("--readback-results", type=Path,
                        help="re-evaluate assertions on this run's native UNO JSON; never reads XLSX caches")
    parser.add_argument("--timeout", type=int, default=3600)
    args = parser.parse_args()
    temporary = tempfile.TemporaryDirectory(prefix="templates-review-") if args.work_dir is None else None
    root = Path(temporary.name) if temporary else args.work_dir.resolve()
    root.mkdir(parents=True, exist_ok=True)
    originals = {name: sha(PRODUCT / "dist" / name) for name in NAMES.values()}
    s = Suite(root, args.reuse_inputs)
    scenarios(s)
    if args.only:
        s.requests = [r for r in s.requests if r["id"].startswith(args.only)]
        s.checks = [c for c in s.checks if c[0].startswith(args.only)]
    started = time.monotonic()
    office_results = (json.loads(args.readback_results.read_text()) if args.readback_results
                     else run_office(args, root, s.requests))
    by_id = {r["id"]: r for r in office_results}
    tests = []
    for id, title, check in s.checks:
        r = by_id.get(id, {"error": "missing scenario result"})
        try:
            if "error" in r:
                tests.append({"scenario": id, "test": title, "status": "BLOCKED", "evidence": r["error"]})
                continue
            ok, evidence = check(r)
            tests.append({"scenario": id, "test": title, "status": "PASS" if ok else "FAIL", "evidence": evidence})
        except Exception as exc:
            tests.append({"scenario": id, "test": title, "status": "FAIL", "evidence": "Native result cannot satisfy assertion: " + str(exc)})
    unchanged = all(sha(PRODUCT / "dist" / name) == digest for name, digest in originals.items())
    tests.append({"scenario": "safety", "test": "shipped files unchanged", "status": "PASS" if unchanged else "FAIL"})
    result = {"environment": {"suite": "Apache OpenOffice", "soffice": str(args.soffice),
        "source_revision": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=PRODUCT).decode().strip(),
        "version_ini": (args.soffice.parent / "version.ini").read_text() if (args.soffice.parent / "version.ini").exists() else "not available",
        "python": sys.version, "openpyxl": __import__("openpyxl").__version__,
        "office_date_basis": str(date.today()), "source_sha256": originals},
        "elapsed_seconds": time.monotonic() - started,
        "native_elapsed_seconds": sum(r.get("load_calculate_save_read_seconds", 0) for r in office_results),
        "assertions_re_evaluated_from_native_json": bool(args.readback_results), "scenarios": office_results,
        "tests": tests, "summary": {state: sum(t["status"] == state for t in tests) for state in ("PASS", "FAIL", "BLOCKED")}}
    # Complete native values remain in the disposable work directory. Commit bounded
    # assertion evidence and native feature inventory, never megabytes of empty rows.
    result["scenarios"] = [{k: v for k, v in r.items() if k != "data"} for r in office_results]
    args.output.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result["summary"]))
    for test in tests:
        if test["status"] != "PASS":
            print(test["status"], test["scenario"], test["test"], json.dumps(test.get("evidence")))
    if temporary:
        temporary.cleanup()
    return 1 if result["summary"]["FAIL"] or result["summary"]["BLOCKED"] else 0


if __name__ == "__main__":
    sys.exit(main())
