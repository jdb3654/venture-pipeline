"""Cross-app test scenarios for the three manufacturing workbooks.

Each scenario starts from a shipped workbook in dist/, types inputs into the yellow cells, and states the values
the grey cells must show. Expected values come from the independent Python model in this file, never from the
workbook's own formulas.

    python3 scenarios.py make <out_dir>     writes <out_dir>/<scenario>.xlsx and <out_dir>/expected.json

Then let each spreadsheet app recalculate and save the files (see README.md), and run check.py on the results.
"""
import json
import shutil
import sys
from datetime import date, time, timedelta
from pathlib import Path

from openpyxl import load_workbook

DIST = Path(__file__).resolve().parents[2] / "dist"
OEE, DT, PM = "OEE-Tracker.xlsx", "Downtime-Log-and-Pareto.xlsx", "Preventive-Maintenance-Schedule.xlsx"
D0 = date(2026, 9, 1)


# --------------------------------------------------------------------------- independent models
def oee_expect(machines, rows, frm=None, to=None):
    """machines: {name: cycle_sec}; rows: (date, machine, shift, stops, down, total, rejects).
    Returns per-row (check, oee or None) and per-machine / total OEE using only valid rows."""
    out_rows, agg = [], {}
    for d, m, sl, ps, dt, tot, rej in rows:
        nums = all(isinstance(v, (int, float)) for v in (sl, tot)) and all(
            v is None or isinstance(v, (int, float)) for v in (ps, dt, rej))
        ps0, dt0, rej0 = (ps or 0), (dt or 0), (rej or 0)
        if m not in machines:
            check = "Machine not in Setup"
        elif not nums:
            check = "Enter numbers: shift length and total parts are required"
        elif sl <= 0 or ps0 < 0 or dt0 < 0 or tot < 0 or rej0 < 0:
            check = "Values must be 0 or more (shift length above 0)"
        elif rej0 > tot:
            check = "Rejects > total parts"
        elif ps0 >= sl:
            check = "Planned stops >= shift length"
        elif dt0 > sl - ps0:
            check = "Downtime > planned time"
        else:
            check = "OK"
        ok = check == "OK"
        oee = None
        if ok:
            ct = machines[m]
            ppt, run = sl - ps0, sl - ps0 - dt0
            perf = (ct * tot / 60) / run if run > 0 else None
            if perf is not None and perf > 1:
                check = "Check: performance above 100%"
            oee = (ct * (tot - rej0) / 60) / ppt
            if (frm is None or d >= frm) and (to is None or d <= to):
                a = agg.setdefault(m, [0.0, 0.0])
                a[0] += ct * (tot - rej0) / 60
                a[1] += ppt
        out_rows.append((check, oee))
    per = {m: (v[0] / v[1] if v[1] else None) for m, v in agg.items()}
    g = sum(v[0] for v in agg.values())
    p = sum(v[1] for v in agg.values())
    return out_rows, per, (g / p if p else None)


def minutes(start, end, typed):
    if typed is not None:
        return typed if isinstance(typed, (int, float)) and typed >= 0 else None
    if start is None or end is None:
        return None
    a = start.hour * 60 + start.minute
    b = end.hour * 60 + end.minute
    return (b - a) % 1440


def dt_expect(machines, reasons, rows, frm=None, to=None, mfilter=None):
    """rows: (date, machine, start, end, typed, reason). Returns totals the Pareto and By Machine sheets must show."""
    tot = 0
    by_reason, stops_reason, by_machine, stops_machine = {}, {}, {}, {}
    all_tot, all_stops = 0, 0
    for d, m, s, e, typed, r in rows:
        mins = minutes(s, e, typed)
        if mins is None or d is None:
            continue
        in_dates = (frm is None or d >= frm) and (to is None or d <= to)
        if not in_dates:
            continue
        all_tot += mins
        all_stops += 1 if mins > 0 else 0
        by_machine[m] = by_machine.get(m, 0) + mins
        stops_machine[m] = stops_machine.get(m, 0) + (1 if mins > 0 else 0)
        if mfilter is not None and m != mfilter:
            continue
        tot += mins
        if r in reasons:
            by_reason[r] = by_reason.get(r, 0) + mins
            stops_reason[r] = stops_reason.get(r, 0) + (1 if mins > 0 else 0)
    classified = sum(by_reason.values())
    return {"total": tot, "unclassified": tot - classified, "by_reason": by_reason, "stops_reason": stops_reason,
            "by_machine": {m: by_machine.get(m, 0) for m in machines},
            "stops_machine": {m: stops_machine.get(m, 0) for m in machines},
            "all_logged": all_tot, "all_stops": all_stops}


# --------------------------------------------------------------------------- scenario writers
def open_clean(name, clear):
    wb = load_workbook(DIST / name)
    for sheet, cells in clear.items():
        ws = wb[sheet]
        for rng in cells:
            for row in ws[rng]:
                for c in row:
                    c.value = None
    return wb


def oee_scenario(name, machines, rows, frm=None, to=None, target=None):
    wb = open_clean(OEE, {"Setup": ["A5:C24"], "Production Log": ["A5:H504"]})
    st, lg, db = wb["Setup"], wb["Production Log"], wb["Dashboard"]
    for i, (m, ct) in enumerate(machines.items()):
        st.cell(5 + i, 1, m)
        st.cell(5 + i, 2, ct)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):  # date, machine, shift, stops, down, total, rejects -> A,C,D,E,F,G,H
            col = [1, 3, 4, 5, 6, 7, 8][j]
            if v is not None:
                lg.cell(5 + i, col, v)
    if frm:
        db["B3"] = frm
    if to:
        db["B4"] = to
    if target is not None:
        db["B5"] = target
    rows_exp, per, total = oee_expect(machines, rows, frm, to)
    exp = []
    for i, (check, oee) in enumerate(rows_exp):
        exp.append({"sheet": "Production Log", "cell": f"R{5 + i}", "value": check})
        exp.append({"sheet": "Production Log", "cell": f"Q{5 + i}", "value": oee if oee is not None else ""})
    for i, m in enumerate(machines):
        exp.append({"sheet": "Dashboard", "cell": f"K{9 + i}", "value": per.get(m, "")})
    exp.append({"sheet": "Dashboard", "cell": "K29", "value": total if total is not None else ""})
    return wb, exp


def dt_scenario(name, machines, rows, frm=None, to=None, mfilter=None, reasons=None):
    reasons = reasons or {"R-Break": "Breakdown", "R-Setup": "Setup", "R-Material": "Material"}
    wb = open_clean(DT, {"Setup": ["A5:A34", "D5:E34"], "Downtime Log": ["A5:F1004", "H5:H1004", "J5:J1004"]})
    st, lg, pa, bm = wb["Setup"], wb["Downtime Log"], wb["Pareto"], wb["By Machine"]
    for i, m in enumerate(machines):
        st.cell(5 + i, 1, m)
    for i, (r, cat) in enumerate(reasons.items()):
        st.cell(5 + i, 4, r)
        st.cell(5 + i, 5, cat)
    for i, (d, m, s, e, typed, r) in enumerate(rows):
        for col, v in zip((1, 3, 4, 5, 6, 8), (d, m, s, e, typed, r)):
            if v is not None:
                lg.cell(5 + i, col, v)
    if frm:
        pa["B3"] = frm
    if to:
        pa["B4"] = to
    if mfilter:
        pa["B5"] = mfilter
    x = dt_expect(machines, reasons, rows, frm, to, mfilter)
    exp = [{"sheet": "Pareto", "cell": "B6", "value": x["total"]},
           {"sheet": "Pareto", "cell": "B7", "value": x["unclassified"]}]
    for i, r in enumerate(reasons):  # hidden helper columns hold per-reason minutes and stops in Setup order
        exp.append({"sheet": "Pareto", "cell": f"J{9 + i}", "value": x["by_reason"].get(r, 0)})
        exp.append({"sheet": "Pareto", "cell": f"K{9 + i}", "value": x["stops_reason"].get(r, 0)})
    for i, m in enumerate(machines):
        exp.append({"sheet": "By Machine", "cell": f"B{5 + i}", "value": x["by_machine"][m]})
        exp.append({"sheet": "By Machine", "cell": f"C{5 + i}", "value": x["stops_machine"][m]})
    exp.append({"sheet": "By Machine", "cell": "B36", "value": x["all_logged"]})
    exp.append({"sheet": "By Machine", "cell": "C36", "value": x["all_stops"]})
    return wb, exp


def pm_scenario(name, plan, log, window=7):
    """plan: (asset, task, freq, last_before_offset_days or None); log: (offset_days or None, asset, task).
    Offsets are days before the app's TODAY(), written as =TODAY()-n so results don't depend on the run date."""
    wb = open_clean(PM, {"PM Plan": ["A5:E304", "J5:J304"], "Completion Log": ["A5:E2004"]})
    pl, cl = wb["PM Plan"], wb["Completion Log"]
    pl["I2"] = window
    for i, (a, t, fr, ago) in enumerate(plan):
        r = 5 + i
        for col, v in ((1, a), (2, t), (3, fr)):
            if v is not None:
                pl.cell(r, col, v)
        if ago is not None:
            pl.cell(r, 5, f"=TODAY()-{ago}")
    for i, (ago, a, t) in enumerate(log):
        r = 5 + i
        if ago is not None:
            cl.cell(r, 1, f"=TODAY()-{ago}")
        cl.cell(r, 2, a)
        cl.cell(r, 3, t)
    exp = []
    keys = [(a, t) for a, t, _, _ in plan]
    counts = {"OVERDUE": 0, "Due soon": 0, "OK": 0, "fix": 0}
    for i, (a, t, fr, ago) in enumerate(plan):
        if a is None or t is None:
            status = "Fill in asset and task"
        elif not isinstance(fr, (int, float)) or fr <= 0:
            status = "Check frequency"
        elif keys.count((a, t)) > 1:
            status = "Duplicate row"
        else:
            done = [n for n, la, lt in log if la == a and lt == t and n is not None]
            cands = done + ([ago] if ago is not None else [])
            days_left = (fr - min(cands)) if cands else 0  # never done: due today
            status = "OVERDUE" if days_left < 0 else ("Due soon" if days_left <= window else "OK")
        exp.append({"sheet": "PM Plan", "cell": f"I{5 + i}", "value": status})
        counts[status if status in counts else "fix"] += 1
    exp += [{"sheet": "PM Plan", "cell": "L1", "value": counts["OVERDUE"]},
            {"sheet": "PM Plan", "cell": "L2", "value": counts["Due soon"]},
            {"sheet": "PM Plan", "cell": "L3", "value": counts["OK"]},
            {"sheet": "PM Plan", "cell": "L4", "value": counts["fix"]}]
    for i, (ago, a, t) in enumerate(log):
        msg = ("Enter the date done" if ago is None else
               ("OK" if (a, t) in keys else "No matching asset + task in plan"))
        exp.append({"sheet": "Completion Log", "cell": f"F{5 + i}", "value": msg})
    return wb, exp


def scenarios():
    d = lambda n: D0 + timedelta(days=n)
    yield "oee_examples", oee_scenario, dict(
        machines={"Press 01": 12, "CNC Lathe 02": 95, "Packaging Line A": 2.5},
        rows=[(d(0), "Press 01", 480, 40, 35, 1900, 22), (d(0), "CNC Lathe 02", 480, 40, 60, 210, 4),
              (d(0), "Packaging Line A", 480, 40, 25, 9200, 150), (d(1), "Press 01", 480, 40, 90, 1500, 30),
              (d(1), "CNC Lathe 02", 480, 40, 20, 240, 2), (d(1), "Packaging Line A", 480, 40, 55, 8400, 95)])
    yield "oee_mixed_cycle_times", oee_scenario, dict(  # Codex #24: correct all-machine total is 50%
        machines={"Slow": 60, "Fast": 6},
        rows=[(d(0), "Slow", 60, 0, 0, 60, 0), (d(0), "Fast", 60, 0, 0, 600, 600)])
    yield "oee_bad_inputs", oee_scenario, dict(  # #25: errors flagged and kept out of totals
        machines={"M1": 12},
        rows=[(d(0), "M1", 480, 40, 35, 1900, 22), (d(0), "M1", 480, 40, 35, 1900, 2000),
              (d(0), "M1", 480, 40, 35, "one thousand", 0), (d(0), "M1", 0, 0, 0, 100, 0),
              (d(0), "M1", 480, 40, 500, 100, 0), (d(0), "Nope", 480, 40, 35, 1900, 22),
              (d(0), "M1", 480, 40, -5, 100, 0), (d(0), "M1", 480, 40, 35, 2065, 0)])
    yield "oee_wildcard_names", oee_scenario, dict(  # #18
        machines={"Line*1": 10, "LineX1": 10, "Jam?2": 10},
        rows=[(d(0), "Line*1", 100, 0, 0, 300, 0), (d(0), "LineX1", 100, 0, 0, 600, 0),
              (d(0), "Jam?2", 100, 0, 0, 100, 0)])
    yield "oee_date_filter", oee_scenario, dict(
        machines={"M1": 12}, frm=d(1), to=d(1),
        rows=[(d(0), "M1", 480, 40, 35, 1900, 22), (d(1), "M1", 480, 40, 90, 1500, 30),
              (d(2), "M1", 480, 40, 10, 2000, 0)])
    yield "down_examples", dt_scenario, dict(
        machines=["Press 01", "CNC Lathe 02", "Packaging Line A"],
        rows=[(d(0), "Press 01", time(7, 15), time(7, 50), None, "R-Setup"),
              (d(0), "Press 01", time(10, 5), time(10, 20), None, "R-Material"),
              (d(0), "CNC Lathe 02", time(9, 0), time(10, 0), None, "R-Break"),
              (d(0), "Packaging Line A", time(23, 40), time(0, 25), None, "R-Break"),
              (d(1), "Press 01", None, None, 55, "R-Break"),
              (d(1), "Packaging Line A", time(13, 0), time(13, 30), None, "R-Setup")])
    yield "down_machine_filter", dt_scenario, dict(
        machines=["M1", "M2"], mfilter="M2",
        rows=[(d(0), "M1", None, None, 30, "R-Break"), (d(0), "M2", None, None, 20, "R-Break"),
              (d(0), "M2", None, None, 5, "R-Setup")])
    yield "down_unknown_negative_zero", dt_scenario, dict(  # #19 #28 #29
        machines=["M1"],
        rows=[(d(0), "M1", None, None, 90, "TYPO"), (d(0), "M1", None, None, 10, "R-Break"),
              (d(0), "M1", None, None, -20, "R-Setup"), (d(0), "M1", None, None, 0, "R-Break"),
              (d(0), "M1", time(8, 0), None, None, "R-Break"), (d(0), "", None, None, 15, "R-Material")])
    yield "down_wildcard_names", dt_scenario, dict(  # #18
        machines=["Line*1", "LineX1"], reasons={"Jam*1": "Breakdown", "JamX1": "Breakdown", "Hold~2": "Quality"},
        rows=[(d(0), "Line*1", None, None, 10, "Jam*1"), (d(0), "LineX1", None, None, 90, "JamX1"),
              (d(0), "LineX1", None, None, 7, "Hold~2")])
    yield "down_dates", dt_scenario, dict(
        machines=["M1"], frm=d(1),
        rows=[(d(0), "M1", None, None, 30, "R-Break"), (d(1), "M1", None, None, 20, "R-Break"),
              (None, "M1", None, None, 99, "R-Break")])
    yield "pm_rules", pm_scenario, dict(  # #20 #30 #31 #18
        plan=[("A1", "T1", 7, None), ("A2", "T2", 30, 40), ("A3", "T3", 0, None), ("A4", "T4", None, None),
              ("A5", "T5", -7, None), ("A6", "T6", 7, None), ("A6", "T6", 30, None), ("Line*1", "Lube", 14, 13),
              ("LineX1", "Lube", 14, None), ("A7", "T7", 90, 100)],
        log=[(3, "A1", "T1"), (1, "A1", "T1"), (None, "A2", "T2"), (5, "A9", "T9"), (2, "LineX1", "Lube"),
             (60, "A7", "T7")])


def make(out):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    expected = {}
    for name, fn, kw in scenarios():
        wb, exp = fn(name, **kw)
        wb.save(out / f"{name}.xlsx")
        expected[name] = exp
    (out / "expected.json").write_text(json.dumps(expected, indent=1, default=str))
    print(f"wrote {len(expected)} scenarios to {out}")


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] != "make":
        sys.exit(__doc__)
    make(sys.argv[2])
