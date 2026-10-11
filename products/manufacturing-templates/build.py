"""Build the manufacturing template bundle (three .xlsx workbooks).

Run: python3 build.py   (writes into ./dist)
Then recalculate each file with LibreOffice before checking values.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import TextAxis
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent / "dist"
FONT = "Arial"
NAVY = "1F3A5F"
INPUT = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="F2F2F2")
HEAD = PatternFill("solid", fgColor=NAVY)
RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6EFCE")
GREY_FIX = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
VERSION = "Version 1.0"


def f(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(16, True, NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(9, italic=True, color="595959")


def headers(ws, row, names, widths=None):
    for i, name in enumerate(names, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = f(10, True, "FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def style_block(ws, first_row, last_row, cols, fill, fmt=None):
    for r in range(first_row, last_row + 1):
        for col in cols:
            c = ws.cell(row=r, column=col)
            c.fill = fill
            c.font = f()
            c.border = BOX
            if fmt:
                c.number_format = fmt


def fmt_cols(ws, first_row, last_row, col_fmts):
    for col, fmt in col_fmts.items():
        for r in range(first_row, last_row + 1):
            ws.cell(row=r, column=col).number_format = fmt


def how_to(ws, name, lines):
    title(ws, name, VERSION)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 100
    r = 4
    for kind, text in lines:
        c = ws.cell(row=r, column=2, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "h":
            c.font = f(12, True, NAVY)
            r += 1
        elif kind == "input":
            c.font = f()
            c.fill = INPUT
            c.border = BOX
        elif kind == "calc":
            c.font = f()
            c.fill = CALC
            c.border = BOX
        else:
            c.font = f()
        r += 1
    ws.sheet_view.showGridLines = False


def traffic_light(ws, rng, target_ref):
    first = rng.split(":")[0].replace("$", "")
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND(ISNUMBER({first}),{first}<0.6)'], fill=RED, stopIfTrue=True))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND(ISNUMBER({first}),{first}>={target_ref})'], fill=GREEN))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND(ISNUMBER({first}),{first}>=0.6,{first}<{target_ref})'], fill=AMBER))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND(ISNUMBER({first}),{first}<0.6)'], fill=RED))


# ---------------------------------------------------------------- OEE tracker
def build_oee():
    wb = Workbook()
    how = wb.active
    how.title = "How to use"
    setup = wb.create_sheet("Setup")
    log = wb.create_sheet("Production Log")
    dash = wb.create_sheet("Dashboard")
    LOG_LAST = 504
    MACH_LAST = 24

    how_to(how, "OEE Tracker", [
        ("h", "What this does"),
        ("t", "Log each shift's production per machine. The workbook calculates Availability, Performance, "
              "Quality and OEE for every shift and summarizes them by machine on the Dashboard, for any date range."),
        ("h", "Colour key"),
        ("input", "Yellow cells are for you to fill in."),
        ("calc", "Grey cells are calculated. Don't type over them."),
        ("h", "Steps"),
        ("t", "1. Setup: list your machines and each one's ideal cycle time (the fastest sustainable time to make one "
              "part, in seconds). For a line, use the time per finished unit at rated speed."),
        ("t", "2. Production Log: add one row per machine per shift. Planned stops are scheduled time the machine is "
              "not expected to run (breaks, planned meetings). Unplanned downtime is everything else it was stopped."),
        ("t", "3. Dashboard: optionally enter a From and To date. Leave them blank to include everything."),
        ("t", "4. The rows marked EXAMPLE show the expected format. Before you start, select their yellow cells and press Delete. Don't delete whole rows."),
        ("h", "How OEE is calculated"),
        ("t", "Planned production time = shift length - planned stops.   Run time = planned production time - unplanned downtime."),
        ("t", "Availability = run time / planned production time."),
        ("t", "Performance = (ideal cycle time x total parts) / run time."),
        ("t", "Quality = good parts / total parts.   OEE = Availability x Performance x Quality."),
        ("t", "Dashboard OEE = total ideal time for good parts / total planned production time. Times and counts are "
              "summed first, so long and short shifts, and machines with different cycle times, are weighted correctly."),
        ("h", "Check column"),
        ("t", "The Check column on the log flags common entry mistakes: a machine missing from Setup, rejects higher than "
              "total parts, downtime longer than planned time, missing or non-number entries. Rows with an error are left "
              "out of the Dashboard. 'Check: performance above 100%' is a warning only (usually a cycle time that is "
              "too slow in Setup, or a miscounted total); that row still counts."),
        ("h", "Target"),
        ("t", "85% OEE is a commonly cited 'world-class' benchmark. Set your own target on the Dashboard; many plants "
              "start well below that, and improvement against your own baseline matters more."),
    ])

    # Setup
    title(setup, "Setup", "List each machine once. Names must match exactly what you pick in the log.")
    headers(setup, 4, ["Machine", "Ideal cycle time\n(seconds per part)", "Notes", "Check"], [26, 20, 40, 36])
    style_block(setup, 5, MACH_LAST, [1, 2, 3], INPUT)
    style_block(setup, 5, MACH_LAST, [4], CALC)
    fmt_cols(setup, 5, MACH_LAST, {2: "0.0"})
    for r in range(5, MACH_LAST + 1):
        setup[f"D{r}"] = f'=IF(A{r}="","",IF(SUMPRODUCT(--($A$5:A{r}=A{r}))>1,"Duplicate: only the first row is used",""))'
    for i, (m, ct, note) in enumerate([
        ("Press 01", 12, "EXAMPLE"),
        ("CNC Lathe 02", 95, "EXAMPLE"),
        ("Packaging Line A", 2.5, "EXAMPLE"),
    ]):
        setup.cell(row=5 + i, column=1, value=m)
        setup.cell(row=5 + i, column=2, value=ct)
        setup.cell(row=5 + i, column=3, value=note)
    setup.freeze_panes = "A5"

    # Production log
    title(log, "Production Log", "One row per machine per shift. Fill in the yellow columns; grey columns calculate.")
    cols = ["Date", "Shift", "Machine", "Shift length\n(min)", "Planned stops\n(min)", "Unplanned\ndowntime (min)",
            "Total parts", "Rejected\nparts", "Planned production\ntime (min)", "Run time\n(min)",
            "Ideal cycle\ntime (sec)", "Ideal run time\n(min)", "Good parts", "Availability", "Performance",
            "Quality", "OEE", "Check", "Ideal time for\ngood parts (min)", "Row used\nin totals"]
    headers(log, 4, cols, [11, 8, 20, 11, 11, 12, 10, 10, 15, 10, 11, 13, 10, 12, 12, 10, 10, 40, 15, 10])
    style_block(log, 5, LOG_LAST, range(1, 9), INPUT)
    style_block(log, 5, LOG_LAST, range(9, 21), CALC)
    fmt_cols(log, 5, LOG_LAST, {1: "yyyy-mm-dd", 9: "0", 10: "0", 11: "0.0", 12: "0.0", 13: "0",
                                 14: "0.0%", 15: "0.0%", 16: "0.0%", 17: "0.0%", 19: "0.0"})
    S = "Setup!$A$5:$A$" + str(MACH_LAST)
    CT = "Setup!$B$5:$B$" + str(MACH_LAST)
    for r in range(5, LOG_LAST + 1):
        # exact-match lookups (SUMPRODUCT) avoid Excel wildcards and IFERROR, which OpenOffice 4.1 lacks
        found = f'SUMPRODUCT(--({S}=$C{r}))>0'
        far = f'SUMPRODUCT(MAX(({S}=$C{r})*({MACH_LAST + 1}-ROW({S}))))'  # first matching Setup row wins
        log[f"K{r}"] = f'=IF($C{r}="","",IF({found},INDEX({CT},{MACH_LAST - 3}-{far}),""))'
        log[f"T{r}"] = (f'=IF($C{r}="",FALSE,AND({found},ISNUMBER(A{r}),ISNUMBER(K{r}),N(K{r})>0,ISNUMBER(D{r}),N(D{r})>0,'
                        f'ISNUMBER(G{r}),N(G{r})>=0,OR(E{r}="",ISNUMBER(E{r})),OR(F{r}="",ISNUMBER(F{r})),'
                        f'OR(H{r}="",ISNUMBER(H{r})),N(E{r})>=0,N(F{r})>=0,N(H{r})>=0,N(H{r})<=N(G{r}),'
                        f'N(E{r})<N(D{r}),N(F{r})<=N(D{r})-N(E{r})))')
        ok = f'$T{r}'
        log[f"I{r}"] = f'=IF({ok},D{r}-N(E{r}),"")'
        log[f"J{r}"] = f'=IF({ok},I{r}-N(F{r}),"")'
        log[f"L{r}"] = f'=IF({ok},K{r}*G{r}/60,"")'
        log[f"M{r}"] = f'=IF({ok},G{r}-N(H{r}),"")'
        log[f"S{r}"] = f'=IF({ok},K{r}*M{r}/60,"")'
        log[f"N{r}"] = f'=IF({ok},J{r}/I{r},"")'
        log[f"O{r}"] = f'=IF(AND({ok},N(J{r})>0),L{r}/J{r},"")'
        log[f"P{r}"] = f'=IF(AND({ok},N(G{r})>0),M{r}/G{r},"")'
        log[f"Q{r}"] = f'=IF({ok},S{r}/I{r},"")'
        log[f"R{r}"] = (f'=IF($C{r}="","",IF(NOT({found}),"Machine not in Setup",'
                        f'IF(OR(NOT(ISNUMBER(K{r})),N(K{r})<=0),"Set an ideal cycle time in Setup",'
                        f'IF(NOT(ISNUMBER(A{r})),"Enter a date",'
                        f'IF(OR(NOT(ISNUMBER(D{r})),NOT(ISNUMBER(G{r})),AND(E{r}<>"",NOT(ISNUMBER(E{r}))),'
                        f'AND(F{r}<>"",NOT(ISNUMBER(F{r}))),AND(H{r}<>"",NOT(ISNUMBER(H{r})))),'
                        f'"Enter numbers: shift length and total parts are required",'
                        f'IF(OR(N(D{r})<=0,N(E{r})<0,N(F{r})<0,N(G{r})<0,N(H{r})<0),"Values must be 0 or more (shift length above 0)",'
                        f'IF(N(H{r})>N(G{r}),"Rejects > total parts",'
                        f'IF(N(E{r})>=N(D{r}),"Planned stops >= shift length",'
                        f'IF(N(F{r})>N(D{r})-N(E{r}),"Downtime > planned time",'
                        f'IF(AND(ISNUMBER(O{r}),O{r}>1),"Check: performance above 100%","OK"))))))))))')
    examples = [
        ("2026-09-01", "Day", "Press 01", 480, 40, 35, 1900, 22),
        ("2026-09-01", "Day", "CNC Lathe 02", 480, 40, 60, 210, 4),
        ("2026-09-01", "Day", "Packaging Line A", 480, 40, 25, 9200, 150),
        ("2026-09-02", "Day", "Press 01", 480, 40, 90, 1500, 30),
        ("2026-09-02", "Day", "CNC Lathe 02", 480, 40, 20, 240, 2),
        ("2026-09-02", "Day", "Packaging Line A", 480, 40, 55, 8400, 95),
    ]
    from datetime import date
    for i, ex in enumerate(examples):
        r = 5 + i
        y, m, d = map(int, ex[0].split("-"))
        log[f"A{r}"] = date(y, m, d)
        for col, val in zip("BCDEFGH", ex[1:]):
            log[f"{col}{r}"] = val
    dv = DataValidation(type="list", formula1=f"=Setup!$A$5:$A${MACH_LAST}", allow_blank=True)
    dv.error = "Pick a machine listed on the Setup sheet."
    log.add_data_validation(dv)
    dv.add(f"C5:C{LOG_LAST}")
    nonneg = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    nonneg.error = "Enter a number of zero or more."
    nonneg.showErrorMessage = True
    log.add_data_validation(nonneg)
    nonneg.add(f"D5:H{LOG_LAST}")
    # Google Sheets ignores colour rules that point at another sheet, so keep a hidden local copy of the target.
    log["U1"] = "=Dashboard!$B$5"
    log.column_dimensions["U"].hidden = True
    traffic_light(log, f"Q5:Q{LOG_LAST}", "$U$1")
    log.freeze_panes = "D5"

    # Dashboard
    title(dash, "Dashboard", "Summary by machine for the date range below. Leave dates blank to include all rows.")
    labels = [("From date", 3), ("To date", 4), ("OEE target", 5)]
    for text, r in labels:
        dash[f"A{r}"] = text
        dash[f"A{r}"].font = f(10, True)
        dash[f"B{r}"].fill = INPUT
        dash[f"B{r}"].border = BOX
        dash[f"B{r}"].font = f()
    dash["B3"].number_format = "yyyy-mm-dd"
    dash["B4"].number_format = "yyyy-mm-dd"
    dash["B5"] = 0.85
    dash["B5"].number_format = "0%"
    tv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=False)
    tv.error = "Enter a target between 0% and 100%."
    dash.add_data_validation(tv)
    tv.add("B5")
    dash["D3"] = '=IF(B3="",DATE(1900,1,1),B3)'
    dash["D4"] = '=IF(B4="",DATE(2999,12,31),B4)'
    for c in ("D3", "D4"):
        dash[c].number_format = "yyyy-mm-dd"
        dash[c].font = f(8, color="FFFFFF")
    cols = ["Machine", "Planned\ntime (min)", "Run time\n(min)", "Unplanned\ndowntime (min)", "Total parts",
            "Good parts", "Ideal run\ntime (min)", "Availability", "Performance", "Quality", "OEE", "OEE vs\ntarget"]
    headers(dash, 8, cols, [22, 12, 12, 14, 12, 12, 12, 12, 12, 11, 10, 11])
    first, last = 9, 9 + (MACH_LAST - 5)
    total = last + 1
    L = "'Production Log'!"
    rng = lambda c: f"{L}${c}$5:${c}${LOG_LAST}"
    src = {2: "I", 3: "J", 4: "F", 5: "G", 6: "M", 7: "L", 13: "S"}
    for r in range(first, last + 1):
        sr = 5 + (r - first)
        dash[f"A{r}"] = (f'=IF(Setup!A{sr}="","",IF(SUMPRODUCT(--(Setup!$A$4:A{sr - 1}=Setup!A{sr}))>0,"",'
                         f'Setup!A{sr}))')
        sel = f'({rng("C")}=$A{r})*({rng("T")}=TRUE)*({rng("A")}>=$D$3)*({rng("A")}<=$D$4)'
        for col, lc in src.items():
            dash.cell(row=r, column=col).value = f'=IF($A{r}="","",SUMPRODUCT({sel},{rng(lc)}))'
        dash[f"H{r}"] = f'=IF($A{r}="","",IF(B{r}=0,"",C{r}/B{r}))'
        dash[f"I{r}"] = f'=IF($A{r}="","",IF(C{r}=0,"",G{r}/C{r}))'
        dash[f"J{r}"] = f'=IF($A{r}="","",IF(G{r}=0,"",M{r}/G{r}))'
        dash[f"K{r}"] = f'=IF($A{r}="","",IF(B{r}=0,"",M{r}/B{r}))'
        dash[f"L{r}"] = f'=IF(K{r}="","",K{r}-$B$5)'
    dash[f"A{total}"] = "All machines"
    for col in list(range(2, 8)) + [13]:
        cl = get_column_letter(col)
        dash[f"{cl}{total}"] = f"=SUM({cl}{first}:{cl}{last})"
    r = total
    dash[f"H{r}"] = f'=IF(B{r}=0,"",C{r}/B{r})'
    dash[f"I{r}"] = f'=IF(C{r}=0,"",G{r}/C{r})'
    dash[f"J{r}"] = f'=IF(G{r}=0,"",M{r}/G{r})'
    dash[f"K{r}"] = f'=IF(B{r}=0,"",M{r}/B{r})'
    dash[f"L{r}"] = f'=IF(K{r}="","",K{r}-$B$5)'
    dash[f"M{first - 1}"] = "Ideal time for good parts (min)"
    dash[f"M{first - 1}"].font = f(8, True, "808080")
    dash.column_dimensions["M"].hidden = True
    dash["A6"] = ("Log rows whose Check shows an error are left out of these totals. "
                  "The 'performance above 100%' warning still counts.")
    dash["A6"].font = f(9, italic=True, color="595959")
    style_block(dash, first, total, range(1, 13), CALC)
    for col in range(1, 13):
        dash.cell(row=total, column=col).font = f(10, True)
    fmt_cols(dash, first, total, {2: "#,##0", 3: "#,##0", 4: "#,##0", 5: "#,##0", 6: "#,##0", 7: "#,##0",
                                    8: "0.0%", 9: "0.0%", 10: "0.0%", 11: "0.0%", 12: "+0.0%;-0.0%;0.0%"})
    traffic_light(dash, f"K{first}:K{total}", "$B$5")

    ch = BarChart()
    ch.type = "col"
    ch.title = "OEE by machine"
    ch.y_axis.title = "OEE"
    ch.y_axis.numFmt = "0%"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 1
    ch.legend = None
    ch.add_data(Reference(dash, min_col=11, min_row=8, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=1, min_row=first, max_row=last))
    ch.height, ch.width = 8, 18
    dash.add_chart(ch, f"A{total + 3}")
    dash.freeze_panes = "B9"

    for ws in (setup, log, dash):
        ws.sheet_view.zoomScale = 100
    wb.active = dash
    return wb


# ---------------------------------------------------------- Downtime / Pareto
REASONS = [
    ("Mechanical breakdown", "Breakdown"),
    ("Electrical / controls fault", "Breakdown"),
    ("Tooling failure", "Breakdown"),
    ("Changeover / setup", "Setup & adjustment"),
    ("Adjustment / tuning", "Setup & adjustment"),
    ("Waiting for material", "Material"),
    ("Waiting for upstream process", "Material"),
    ("Quality hold / inspection", "Quality"),
    ("No operator available", "Labor"),
    ("Planned maintenance", "Planned"),
    ("Meeting / training", "Planned"),
    ("Other (describe in notes)", "Other"),
]


def build_downtime():
    wb = Workbook()
    how = wb.active
    how.title = "How to use"
    setup = wb.create_sheet("Setup")
    log = wb.create_sheet("Downtime Log")
    par = wb.create_sheet("Pareto")
    bym = wb.create_sheet("By Machine")
    LOG_LAST = 1004
    LIST_LAST = 34  # 30 machines / 30 reasons

    how_to(how, "Downtime Log & Pareto", [
        ("h", "What this does"),
        ("t", "Record every stop with a reason. The Pareto sheet ranks reasons by minutes lost so you can see the "
              "few causes behind most of your downtime, for any date range and any machine."),
        ("h", "Colour key"),
        ("input", "Yellow cells are for you to fill in."),
        ("calc", "Grey cells are calculated. Don't type over them."),
        ("h", "Steps"),
        ("t", "1. Setup: list your machines, and edit the reason list to match how your plant talks about stops. "
              "Each reason has a category (Breakdown, Material, and so on) that you can also rename."),
        ("t", "2. Downtime Log: one row per stop. Enter start and end times, or type the minutes directly if you only "
              "know the duration. Stops that cross midnight are handled automatically."),
        ("t", "3. Pareto: choose a date range and, optionally, one machine. Reasons are ranked automatically, with the "
              "cumulative percentage and a chart."),
        ("t", "4. By Machine: total downtime and stop counts per machine for the same date range."),
        ("t", "5. The rows marked EXAMPLE show the expected format. Before you start, select their yellow cells and press Delete. Don't delete whole rows."),
        ("h", "Tips"),
        ("t", "Keep the reason list short (10-15 reasons). If 'Other' becomes one of your top reasons, read its notes "
              "and add the missing reason."),
        ("t", "Pick reasons from the drop-down so spelling stays consistent. A typed reason that isn't on the Setup "
              "list shows 'Not in Setup'; its minutes are counted as Unclassified on the Pareto sheet until you fix it."),
        ("t", "The Check column on the log should say OK. A stop counts when its duration is more than 0 minutes. "
              "Rows without a date, or with minutes that aren't a number of 0 or more, are left out of the totals."),
    ])

    title(setup, "Setup", "Machines and downtime reasons used by the drop-downs.")
    headers(setup, 4, ["Machine", "Check", "", "Downtime reason", "Category", "Check"], [24, 30, 3, 34, 22, 30])
    setup["C4"].fill = PatternFill(None)
    setup["C4"].border = Border()
    style_block(setup, 5, LIST_LAST, [1, 4, 5], INPUT)
    style_block(setup, 5, LIST_LAST, [2, 6], CALC)
    for r in range(5, LIST_LAST + 1):
        setup[f"B{r}"] = f'=IF(A{r}="","",IF(SUMPRODUCT(--($A$5:A{r}=A{r}))>1,"Duplicate: only the first row is used",""))'
        setup[f"F{r}"] = f'=IF(D{r}="","",IF(SUMPRODUCT(--($D$5:D{r}=D{r}))>1,"Duplicate: only the first row is used",""))'
    for i, m in enumerate(["Press 01", "CNC Lathe 02", "Packaging Line A"]):
        setup.cell(row=5 + i, column=1, value=m)
    for i, (reason, cat) in enumerate(REASONS):
        setup.cell(row=5 + i, column=4, value=reason)
        setup.cell(row=5 + i, column=5, value=cat)
    setup.freeze_panes = "A5"

    title(log, "Downtime Log", "One row per stop. Use start/end times, or type minutes if you only know the duration.")
    cols = ["Date", "Shift", "Machine", "Start\n(hh:mm)", "End\n(hh:mm)", "Minutes (if no\nstart/end)",
            "Duration\n(min)", "Reason", "Category", "Notes", "Check"]
    headers(log, 4, cols, [11, 8, 20, 9, 9, 13, 10, 30, 20, 40, 34])
    style_block(log, 5, LOG_LAST, [1, 2, 3, 4, 5, 6, 8, 10], INPUT)
    style_block(log, 5, LOG_LAST, [7, 9, 11], CALC)
    fmt_cols(log, 5, LOG_LAST, {1: "yyyy-mm-dd", 4: "hh:mm", 5: "hh:mm", 6: "0", 7: "0"})
    RS = f"Setup!$D$5:$D${LIST_LAST}"
    MS = f"Setup!$A$5:$A${LIST_LAST}"
    for r in range(5, LOG_LAST + 1):
        # A stop counts when Duration is a number above 0. Bad typed minutes give a blank duration and a Check message.
        log[f"G{r}"] = (f'=IF(F{r}<>"",IF(AND(ISNUMBER(F{r}),N(F{r})>=0),F{r},""),'
                        f'IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),ROUND(MOD(E{r}-D{r},1)*1440,0),""))')
        far = f'SUMPRODUCT(MAX(({RS}=H{r})*({LIST_LAST + 1}-ROW({RS}))))'  # exact match, first Setup row wins
        log[f"I{r}"] = (f'=IF(H{r}="","",IF({far}=0,"Not in Setup",'
                        f'INDEX(Setup!$E$5:$E${LIST_LAST},{LIST_LAST - 3}-{far})))')
        log[f"K{r}"] = (f'=IF(AND(A{r}="",C{r}="",D{r}="",E{r}="",F{r}="",H{r}=""),"",'
                        f'IF(NOT(ISNUMBER(A{r})),"Enter a date",'
                        f'IF(AND(F{r}<>"",OR(NOT(ISNUMBER(F{r})),N(F{r})<0)),"Minutes must be a number of 0 or more",'
                        f'IF(G{r}="","Enter start and end times, or minutes",'
                        f'IF(C{r}="","Pick a machine",IF(SUMPRODUCT(--({MS}=C{r}))=0,"Machine not in Setup",'
                        f'IF(H{r}="","Pick a reason (counted as unclassified)",'
                        f'IF(I{r}="Not in Setup","Reason not in Setup (counted as unclassified)",'
                        f'IF(N(G{r})=0,"0 minutes: not counted as a stop","OK")))))))))')
    from datetime import date, time
    ex = [
        (date(2026, 9, 1), "Day", "Press 01", time(7, 15), time(7, 50), None, "Changeover / setup", "Die change"),
        (date(2026, 9, 1), "Day", "Press 01", time(10, 5), time(10, 20), None, "Waiting for material", ""),
        (date(2026, 9, 1), "Day", "CNC Lathe 02", time(9, 0), time(10, 0), None, "Tooling failure", "Insert broke"),
        (date(2026, 9, 1), "Night", "Packaging Line A", time(23, 40), time(0, 25), None, "Mechanical breakdown",
         "Conveyor jam, crossed midnight"),
        (date(2026, 9, 2), "Day", "Press 01", None, None, 55, "Mechanical breakdown", "Minutes typed directly"),
        (date(2026, 9, 2), "Day", "Packaging Line A", time(13, 0), time(13, 30), None, "Changeover / setup", ""),
        (date(2026, 9, 2), "Day", "CNC Lathe 02", time(14, 10), time(14, 25), None, "Quality hold / inspection", ""),
    ]
    for i, row in enumerate(ex):
        r = 5 + i
        for col, val in zip("ABCDEFHJ", row):
            if val is not None and val != "":
                log[f"{col}{r}"] = val
        log[f"J{r}"] = ("EXAMPLE - " + row[7]) if row[7] else "EXAMPLE"
    dv_m = DataValidation(type="list", formula1=f"=Setup!$A$5:$A${LIST_LAST}", allow_blank=True)
    dv_r = DataValidation(type="list", formula1=f"=Setup!$D$5:$D${LIST_LAST}", allow_blank=True)
    dv_m.error = "Pick a machine listed on the Setup sheet."
    dv_r.error = "Pick a reason listed on the Setup sheet."
    for dv, rng in ((dv_m, f"C5:C{LOG_LAST}"), (dv_r, f"H5:H{LOG_LAST}")):
        log.add_data_validation(dv)
        dv.add(rng)
    mins = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    mins.error = "Enter minutes as a number of 0 or more."
    log.add_data_validation(mins)
    mins.add(f"F5:F{LOG_LAST}")
    log.freeze_panes = "D5"

    # Pareto
    title(par, "Downtime Pareto", "Reasons ranked by minutes lost. Leave filters blank to include everything.")
    for text, r in (("From date", 3), ("To date", 4), ("Machine (blank = all)", 5)):
        par[f"A{r}"] = text
        par[f"A{r}"].font = f(10, True)
        par[f"B{r}"].fill = INPUT
        par[f"B{r}"].border = BOX
        par[f"B{r}"].font = f()
    par["B3"].number_format = "yyyy-mm-dd"
    par["B4"].number_format = "yyyy-mm-dd"
    dvp = DataValidation(type="list", formula1=f"=Setup!$A$5:$A${LIST_LAST}", allow_blank=True)
    par.add_data_validation(dvp)
    dvp.add("B5")
    par["D3"] = '=IF(B3="",DATE(1900,1,1),B3)'
    par["D4"] = '=IF(B4="",DATE(2999,12,31),B4)'
    par["D3"].number_format = par["D4"].number_format = "yyyy-mm-dd"
    for c in ("D3", "D4"):
        par[c].font = f(8, color="FFFFFF")
    par["A6"] = "Total downtime (min)"
    par["A6"].font = f(10, True)
    first, last = 9, 9 + (LIST_LAST - 5)
    D = "'Downtime Log'!"
    dcol = lambda c: f"{D}${c}$5:${c}${LOG_LAST}"
    sel = (f'({dcol("A")}>=$D$3)*({dcol("A")}<=$D$4)*((($B$5="")+({dcol("C")}=$B$5))>0)')
    par["B6"] = f'=SUMPRODUCT({sel},{dcol("G")})'
    par["B6"].number_format = "#,##0"
    par["B6"].font = f(10, True)
    par["A7"] = "Unclassified (min)"
    par["A7"].font = f(10, True)
    par["B7"] = f"=B6-SUM($J${first}:$J${last})"
    par["B7"].number_format = "#,##0"
    par["C7"] = "Reason blank or not on Setup. Check the log."
    par["C7"].font = f(8, italic=True, color="595959")
    headers(par, 8, ["Rank", "Reason", "Minutes", "Stops", "% of total", "Cumulative %"],
            [22, 32, 11, 9, 11, 13])
    for c, w in (("H", 3), ("I", 30), ("J", 11), ("K", 9), ("L", 14)):
        par.column_dimensions[c].width = w
    for col, name in zip("IJKL", ["Helper: reason", "Helper: minutes", "Helper: stops", "Helper: sort key"]):
        par[f"{col}8"] = name
        par[f"{col}8"].font = f(8, True, "808080")
    for r in range(first, last + 1):
        sr = 5 + (r - first)
        k = r - first + 1
        par[f"I{r}"] = (f'=IF(Setup!D{sr}="","",IF(SUMPRODUCT(--(Setup!$D$4:D{sr - 1}=Setup!D{sr}))>0,"",'
                        f'Setup!D{sr}))')
        par[f"J{r}"] = f'=IF(I{r}="",0,SUMPRODUCT(({dcol("H")}=I{r})*{sel},{dcol("G")}))'
        par[f"K{r}"] = (f'=IF(I{r}="",0,SUMPRODUCT(({dcol("H")}=I{r})*{sel}'
                        f'*ISNUMBER({dcol("G")})*({dcol("G")}>0)))')
        par[f"L{r}"] = f'=IF(I{r}="",-1,J{r}+({LIST_LAST}-{k})/1000000)'
        par[f"A{r}"] = k
        pick = f'MATCH(LARGE($L${first}:$L${last},A{r}),$L${first}:$L${last},0)'
        par[f"B{r}"] = f'=IF(INDEX($J${first}:$J${last},{pick})<=0,"",INDEX($I${first}:$I${last},{pick}))'
        par[f"C{r}"] = f'=IF(B{r}="","",INDEX($J${first}:$J${last},{pick}))'
        par[f"D{r}"] = f'=IF(B{r}="","",INDEX($K${first}:$K${last},{pick}))'
        par[f"E{r}"] = f'=IF(OR(B{r}="",$B$6=0),"",C{r}/$B$6)'
        par[f"F{r}"] = f'=IF(OR(B{r}="",$B$6=0),"",SUM($C${first}:C{r})/$B$6)'
        # Chart-only copy of cumulative %: carries the last value forward so empty rows don't plot as 0.
        par[f"M{r}"] = f'=IF(B{r}<>"",F{r},' + ('""' if r == first else f'M{r - 1}') + ')'
        for col in "IJKLM":
            par[f"{col}{r}"].font = f(8, color="808080")
    style_block(par, first, last, range(1, 7), CALC)
    fmt_cols(par, first, last, {3: "#,##0", 4: "0", 5: "0.0%", 6: "0.0%"})
    for r in range(first, last + 1):
        par[f"A{r}"].alignment = Alignment(horizontal="center")

    bar = BarChart()
    bar.type = "col"
    bar.title = "Downtime by reason (top 12)"
    bar.y_axis.title = "Minutes"
    bar.add_data(Reference(par, min_col=3, min_row=8, max_row=first + 11), titles_from_data=True)
    bar.set_categories(Reference(par, min_col=2, min_row=first, max_row=first + 11))
    line = LineChart()
    par["M8"] = "Cumulative %"
    line.add_data(Reference(par, min_col=13, min_row=8, max_row=first + 11), titles_from_data=True)
    line.y_axis.axId = 200
    line.y_axis.title = "Cumulative %"
    line.y_axis.numFmt = "0%"
    line.y_axis.scaling.min = 0
    line.y_axis.scaling.max = 1
    line.y_axis.crosses = "max"
    # Give the line its own hidden category axis, so Excel doesn't print the reason labels twice.
    line.x_axis = TextAxis(axId=500, delete=True, crosses="autoZero")
    line.y_axis.crossAx = 500
    # Tilt the reason labels so long names don't run together.
    bar.x_axis.txPr = RichText(bodyPr=RichTextProperties(rot=-2700000, vert="horz"),
                               p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=800)),
                                            endParaRPr=CharacterProperties(sz=800))])
    bar += line
    bar.visible_cells_only = False  # the line reads the hidden helper column M
    bar.height, bar.width = 9, 20
    for series in line.series:
        series.smooth = False
    for col in "IJKLM":
        par.column_dimensions[col].hidden = True
    par.add_chart(bar, "N9")  # right after the hidden helper columns I:M; Sheets mis-places charts that span hidden columns
    par.freeze_panes = "A9"

    # By machine
    title(bym, "Downtime by Machine", "Uses the date range set on the Pareto sheet (machine filter is ignored here). A stop is any row with more than 0 minutes.")
    headers(bym, 4, ["Machine", "Downtime (min)", "Stops", "Avg min per stop", "% of all downtime"],
            [24, 15, 10, 16, 18])
    mfirst, mlast = 5, 5 + (LIST_LAST - 5)
    dsel = f'({dcol("A")}>=Pareto!$D$3)*({dcol("A")}<=Pareto!$D$4)'
    other, tot = mlast + 1, mlast + 2
    for r in range(mfirst, mlast + 1):
        bym[f"A{r}"] = (f'=IF(Setup!A{r}="","",IF(SUMPRODUCT(--(Setup!$A$4:A{r - 1}=Setup!A{r}))>0,"",'
                        f'Setup!A{r}))')
        bym[f"B{r}"] = f'=IF($A{r}="","",SUMPRODUCT(({dcol("C")}=$A{r})*{dsel},{dcol("G")}))'
        bym[f"C{r}"] = (f'=IF($A{r}="","",SUMPRODUCT(({dcol("C")}=$A{r})*{dsel}'
                        f'*ISNUMBER({dcol("G")})*({dcol("G")}>0)))')
    bym[f"A{other}"] = "No machine / not in Setup"
    bym[f"B{other}"] = f"=B{tot}-SUM(B{mfirst}:B{mlast})"
    bym[f"C{other}"] = f'=SUMPRODUCT({dsel}*ISNUMBER({dcol("G")})*({dcol("G")}>0))-SUM(C{mfirst}:C{mlast})'
    bym[f"A{tot}"] = "All logged downtime"
    bym[f"B{tot}"] = f'=SUMPRODUCT({dsel},{dcol("G")})'
    bym[f"C{tot}"] = f'=SUMPRODUCT({dsel}*ISNUMBER({dcol("G")})*({dcol("G")}>0))'
    for r in range(mfirst, tot + 1):
        bym[f"D{r}"] = f'=IF(OR($A{r}="",N(C{r})=0),"",B{r}/C{r})'
        bym[f"E{r}"] = f'=IF(OR($A{r}="",$B${tot}=0),"",B{r}/$B${tot})'
    style_block(bym, mfirst, tot, range(1, 6), CALC)
    for c in range(1, 6):
        bym.cell(row=tot, column=c).font = f(10, True)
    fmt_cols(bym, mfirst, tot, {2: "#,##0", 3: "0", 4: "0.0", 5: "0.0%"})
    bym.freeze_panes = "A5"
    wb.active = par
    return wb


# ------------------------------------------------------------- PM schedule
def build_pm():
    wb = Workbook()
    how = wb.active
    how.title = "How to use"
    plan = wb.create_sheet("PM Plan")
    done = wb.create_sheet("Completion Log")
    PLAN_LAST = 304
    LOG_LAST = 2004

    how_to(how, "Preventive Maintenance Schedule", [
        ("h", "What this does"),
        ("t", "List every recurring maintenance task with how often it's due. Log completions as they happen, and the "
              "plan shows each task's next due date and whether it is overdue, due soon or OK."),
        ("h", "Colour key"),
        ("input", "Yellow cells are for you to fill in."),
        ("calc", "Grey cells are calculated. Don't type over them."),
        ("h", "Steps"),
        ("t", "1. PM Plan: one row per asset + task, with the frequency in days (7 = every week, 30 = about monthly, "
              "90 = about quarterly, 365 = about yearly). Due dates count fixed days from the last completion. If you know when a task was last done before you started using this "
              "workbook, enter it under 'Last done before log'."),
        ("t", "2. Completion Log: each time a task is done, add a row with the date, the asset and the task, picked from "
              "the drop-downs so they match the plan exactly. The 'Matches plan?' column warns you if they don't."),
        ("t", "3. The plan updates automatically: last completed date, next due date, days left and status."),
        ("t", "4. Change the 'Due soon' window at the top of the PM Plan (default 7 days)."),
        ("t", "5. A task with no completion and no earlier date shows as due today, every day, until you log it or enter a last-done date."),
        ("t", "Status 'Check frequency', 'Duplicate row' or 'Fill in asset and task' means that row needs fixing; the Needs fixing count is at the top."),
        ("t", "6. The rows marked EXAMPLE show the expected format. Before you start, select their yellow cells and press Delete. Don't delete whole rows."),
        ("h", "Tip"),
        ("t", "Sort or filter the PM Plan by Status to get this week's work list. Excel: Data > Filter."),
    ])

    title(plan, "PM Plan", "One row per asset and task.")
    plan["G2"] = "Due soon window (days)"
    plan["G2"].font = f(10, True)
    plan["G2"].alignment = Alignment(horizontal="right")
    plan.merge_cells("G2:H2")
    plan["I2"] = 7
    plan["I2"].fill = INPUT
    plan["I2"].border = BOX
    plan["I2"].font = f()
    plan["K1"] = "Overdue"
    plan["K2"] = "Due soon"
    plan["K3"] = "OK"
    plan["L1"] = f'=COUNTIF($I$5:$I${PLAN_LAST},"OVERDUE")'
    plan["L2"] = f'=COUNTIF($I$5:$I${PLAN_LAST},"Due soon")'
    plan["L3"] = f'=COUNTIF($I$5:$I${PLAN_LAST},"OK")'
    plan["K4"] = "Needs fixing"
    plan["L4"] = (f'=COUNTIF($I$5:$I${PLAN_LAST},"Check frequency")+COUNTIF($I$5:$I${PLAN_LAST},"Duplicate row")'
                  f'+COUNTIF($I$5:$I${PLAN_LAST},"Fill in asset and task")')
    for c in ("K4", "L4"):
        plan[c].font = f(10, True)
        plan[c].fill = GREY_FIX
    for r, fill in ((1, RED), (2, AMBER), (3, GREEN)):
        plan[f"K{r}"].font = f(10, True)
        plan[f"K{r}"].fill = fill
        plan[f"L{r}"].font = f(10, True)
        plan[f"L{r}"].fill = fill
    cols = ["Asset", "Task", "Frequency\n(days)", "Assigned to", "Last done\nbefore log", "Last completed",
            "Next due", "Days until due", "Status", "Notes"]
    headers(plan, 4, cols, [22, 34, 11, 16, 13, 14, 12, 12, 12, 30])
    style_block(plan, 5, PLAN_LAST, [1, 2, 3, 4, 5, 10], INPUT)
    style_block(plan, 5, PLAN_LAST, [6, 7, 8, 9], CALC)
    fmt_cols(plan, 5, PLAN_LAST, {5: "yyyy-mm-dd", 6: "yyyy-mm-dd", 7: "yyyy-mm-dd", 8: "0"})
    C = "'Completion Log'!"
    for r in range(5, PLAN_LAST + 1):
        mx = (f'MAX(N(E{r}),SUMPRODUCT(MAX(({C}$B$5:$B${LOG_LAST}=A{r})*({C}$C$5:$C${LOG_LAST}=B{r})'
              f'*{C}$A$5:$A${LOG_LAST})))')
        plan[f"F{r}"] = f'=IF(OR(A{r}="",B{r}=""),"",IF({mx}=0,"",{mx}))'
        bad_freq = f'OR(NOT(ISNUMBER(C{r})),N(C{r})<=0)'
        dup = f'SUMPRODUCT(($A$5:$A${PLAN_LAST}=A{r})*($B$5:$B${PLAN_LAST}=B{r}))>1'
        plan[f"G{r}"] = f'=IF(OR(A{r}="",B{r}="",{bad_freq}),"",IF(F{r}="",TODAY(),F{r}+C{r}))'
        plan[f"H{r}"] = f'=IF(G{r}="","",G{r}-TODAY())'
        plan[f"I{r}"] = (f'=IF(AND(A{r}="",B{r}=""),"",IF(OR(A{r}="",B{r}=""),"Fill in asset and task",'
                         f'IF({bad_freq},"Check frequency",IF({dup},"Duplicate row",'
                         f'IF(H{r}<0,"OVERDUE",IF(H{r}<=$I$2,"Due soon","OK"))))))')
    rng = f"I5:I{PLAN_LAST}"
    plan.conditional_formatting.add(rng, FormulaRule(formula=['I5="OVERDUE"'], fill=RED))
    plan.conditional_formatting.add(rng, FormulaRule(formula=['I5="Due soon"'], fill=AMBER))
    plan.conditional_formatting.add(rng, FormulaRule(formula=['I5="OK"'], fill=GREEN))
    plan.conditional_formatting.add(rng, FormulaRule(
        formula=['OR(I5="Check frequency",I5="Duplicate row",I5="Fill in asset and task")'], fill=GREY_FIX))
    fq = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
    fq.error = "Enter the frequency as a whole number of days, 1 or more."
    plan.add_data_validation(fq)
    fq.add(f"C5:C{PLAN_LAST}")
    ex = [
        ("Air compressor 1", "Change oil filter", 90, "Maintenance", 100),
        ("Press 01", "Lubricate slides and gibs", 7, "Operator", 3),
        ("Forklift 2", "Inspect forks, chains and hydraulics", 30, "Maintenance", 27),
        ("CNC Lathe 02", "Clean chip conveyor and coolant filter", 14, "Operator", None),
        ("Dust collector", "Replace filter cartridges", 180, "Maintenance", 20),
    ]
    for i, (a, t, fr, who, ago) in enumerate(ex):
        r = 5 + i
        plan[f"A{r}"], plan[f"B{r}"], plan[f"C{r}"], plan[f"D{r}"] = a, t, fr, who
        if ago is not None:
            plan[f"E{r}"] = f"=TODAY()-{ago}"
        plan[f"J{r}"] = "EXAMPLE"
    plan.freeze_panes = "C5"

    title(done, "Completion Log", "Add a row each time a task is completed. Pick asset and task from the drop-downs.")
    headers(done, 4, ["Date done", "Asset", "Task", "Done by", "Notes", "Matches plan?"], [12, 22, 34, 16, 34, 26])
    style_block(done, 5, LOG_LAST, [1, 2, 3, 4, 5], INPUT)
    style_block(done, 5, LOG_LAST, [6], CALC)
    fmt_cols(done, 5, LOG_LAST, {1: "yyyy-mm-dd"})
    P = "'PM Plan'!"
    for r in range(5, LOG_LAST + 1):
        done[f"F{r}"] = (f'=IF(AND(A{r}="",B{r}="",C{r}=""),"",IF(NOT(ISNUMBER(A{r})),"Enter the date done",'
                         f'IF(OR(B{r}="",C{r}=""),"Fill in asset and task",IF(SUMPRODUCT(({P}$A$5:$A${PLAN_LAST}=B{r})*({P}$B$5:$B${PLAN_LAST}=C{r}))=0,'
                         f'"No matching asset + task in plan","OK"))))')
    done["A5"] = "=TODAY()-2"
    done["B5"] = "CNC Lathe 02"
    done["C5"] = "Clean chip conveyor and coolant filter"
    done["D5"] = "J. Smith"
    done["E5"] = "EXAMPLE"
    dv_a = DataValidation(type="list", formula1=f"={P}$A$5:$A${PLAN_LAST}", allow_blank=True)
    dv_t = DataValidation(type="list", formula1=f"={P}$B$5:$B${PLAN_LAST}", allow_blank=True)
    for dv, rng in ((dv_a, f"B5:B{LOG_LAST}"), (dv_t, f"C5:C{LOG_LAST}")):
        dv.showErrorMessage = False  # warn via the 'Matches plan?' column instead of blocking
        done.add_data_validation(dv)
        dv.add(rng)
    dd = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True)
    dd.error = "Enter the date the task was done."
    done.add_data_validation(dd)
    dd.add(f"A5:A{LOG_LAST}")
    done.freeze_panes = "A5"
    wb.active = plan
    return wb


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for name, fn in (("OEE-Tracker.xlsx", build_oee), ("Downtime-Log-and-Pareto.xlsx", build_downtime),
                     ("Preventive-Maintenance-Schedule.xlsx", build_pm)):
        wb = fn()
        for ws in wb.worksheets:
            ws.sheet_properties.tabColor = NAVY if ws.title == "How to use" else None
        wb.save(OUT / name)
        print("wrote", OUT / name)
